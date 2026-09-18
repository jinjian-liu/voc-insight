import csv
import io
from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.database import get_db
from app.analysis_tasks import analyze_feedback_task
from app.dependencies import require_runtime_session
from app.models import Feedback, FeedbackAnalysis
from app.schemas import AnalysisUpdate, AnalysisView, FeedbackCreate, FeedbackPage, FeedbackView, ImportResult
from app.session_store import RuntimeSession


router = APIRouter(prefix="/feedbacks", tags=["feedbacks"])


@router.post("", response_model=FeedbackView, status_code=status.HTTP_201_CREATED)
def create_feedback(
    payload: FeedbackCreate,
    db: Session = Depends(get_db),
    runtime: RuntimeSession = Depends(require_runtime_session),
) -> Feedback:
    feedback = Feedback(**payload.model_dump(), created_by_name=runtime.operator_name)
    db.add(feedback)
    db.commit()
    db.refresh(feedback)
    return feedback


@router.get("", response_model=FeedbackPage)
def list_feedbacks(
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    keyword: str | None = Query(default=None, max_length=100),
    source: str | None = Query(default=None, max_length=50),
    db: Session = Depends(get_db),
    _: RuntimeSession = Depends(require_runtime_session),
) -> FeedbackPage:
    conditions = [Feedback.deleted_at.is_(None)]
    if keyword:
        pattern = f"%{keyword.strip()}%"
        conditions.append(or_(Feedback.content.ilike(pattern), Feedback.product_module.ilike(pattern)))
    if source:
        conditions.append(Feedback.source == source)

    total = db.scalar(select(func.count(Feedback.id)).where(*conditions)) or 0
    items = db.scalars(
        select(Feedback)
        .where(*conditions)
        .order_by(Feedback.feedback_time.desc(), Feedback.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return FeedbackPage(items=list(items), total=total, page=page, page_size=page_size)


def _normalize_row(row: dict[str, object]) -> dict[str, object]:
    aliases = {
        "content": ["content", "反馈内容", "内容"],
        "source": ["source", "来源", "反馈来源"],
        "feedback_time": ["feedback_time", "反馈时间", "时间"],
        "product_module": ["product_module", "产品模块", "模块"],
        "external_id": ["external_id", "外部记录编号", "外部编号"],
        "note": ["note", "备注"],
    }
    normalized: dict[str, object] = {}
    for target, names in aliases.items():
        normalized[target] = next((row.get(name) for name in names if row.get(name) not in (None, "")), None)
    return normalized


def _parse_time(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if value:
        text = str(value).strip()
        for parser in (datetime.fromisoformat, lambda item: datetime.strptime(item, "%Y-%m-%d %H:%M:%S"), lambda item: datetime.strptime(item, "%Y-%m-%d")):
            try:
                return parser(text)
            except ValueError:
                continue
    return datetime.now(timezone.utc)


@router.post("/import", response_model=ImportResult)
async def import_feedbacks(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    runtime: RuntimeSession = Depends(require_runtime_session),
) -> ImportResult:
    max_bytes = 10 * 1024 * 1024
    content = await file.read(max_bytes + 1)
    await file.close()
    if len(content) > max_bytes:
        raise HTTPException(status_code=413, detail="上传文件不能超过 10 MB")
    suffix = Path(file.filename or "").suffix.lower()
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            rows = list(csv.DictReader(io.StringIO(text)))
        elif suffix == ".xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = [str(item).strip() if item is not None else "" for item in next(values)]
            rows = [dict(zip(headers, row)) for row in values]
            workbook.close()
        else:
            raise HTTPException(status_code=400, detail="仅支持 .csv 和 .xlsx 文件")
    except (UnicodeDecodeError, StopIteration, ValueError) as exc:
        raise HTTPException(status_code=400, detail=f"文件解析失败：{exc}")

    if len(rows) > 2_000:
        raise HTTPException(status_code=400, detail="单次最多导入 2,000 条反馈")
    success = 0
    errors: list[str] = []
    for index, source_row in enumerate(rows, start=2):
        row = _normalize_row(source_row)
        text = str(row.get("content") or "").strip()
        if not text:
            errors.append(f"第 {index} 行：反馈内容为空")
            continue
        external_id = str(row["external_id"]).strip() if row.get("external_id") else None
        if external_id and db.scalar(select(Feedback.id).where(Feedback.external_id == external_id, Feedback.deleted_at.is_(None))):
            errors.append(f"第 {index} 行：外部记录编号重复")
            continue
        db.add(Feedback(
            content=text[:10_000],
            source=str(row.get("source") or "文件导入").strip()[:50],
            feedback_time=_parse_time(row.get("feedback_time")),
            product_module=(str(row["product_module"]).strip()[:100] if row.get("product_module") else None),
            external_id=external_id[:100] if external_id else None,
            note=(str(row["note"]).strip()[:2_000] if row.get("note") else None),
            created_by_name=runtime.operator_name,
        ))
        success += 1
    db.commit()
    return ImportResult(success=success, failed=len(rows) - success, errors=errors[:50])


@router.get("/{feedback_id}", response_model=FeedbackView)
def get_feedback(
    feedback_id: int,
    db: Session = Depends(get_db),
    _: RuntimeSession = Depends(require_runtime_session),
) -> Feedback:
    feedback = db.get(Feedback, feedback_id)
    if feedback is None or feedback.deleted_at is not None:
        raise HTTPException(status_code=404, detail="反馈不存在")
    return feedback


@router.post("/{feedback_id}/analyze", status_code=status.HTTP_202_ACCEPTED)
def start_analysis(
    feedback_id: int,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    runtime: RuntimeSession = Depends(require_runtime_session),
) -> dict[str, str]:
    feedback = db.get(Feedback, feedback_id)
    if feedback is None or feedback.deleted_at is not None:
        raise HTTPException(status_code=404, detail="反馈不存在")
    if feedback.analysis_status == "analyzing":
        return {"status": "analyzing"}
    feedback.analysis_status = "analyzing"
    db.commit()
    background_tasks.add_task(analyze_feedback_task, feedback_id, runtime.api_key)
    return {"status": "analyzing"}


@router.put("/{feedback_id}/analysis", response_model=AnalysisView)
def update_analysis(
    feedback_id: int,
    payload: AnalysisUpdate,
    db: Session = Depends(get_db),
    runtime: RuntimeSession = Depends(require_runtime_session),
) -> FeedbackAnalysis:
    analysis = db.scalar(
        select(FeedbackAnalysis).where(FeedbackAnalysis.feedback_id == feedback_id)
    )
    if analysis is None:
        raise HTTPException(status_code=404, detail="分析结果不存在")
    for key, value in payload.model_dump(exclude={"confirmed"}).items():
        setattr(analysis, key, value)
    analysis.confirmed = payload.confirmed
    analysis.confirmed_by_name = runtime.operator_name if payload.confirmed else None
    analysis.confirmed_at = datetime.now(timezone.utc) if payload.confirmed else None
    db.commit()
    db.refresh(analysis)
    return analysis


@router.delete("/{feedback_id}", status_code=204)
def delete_feedback(
    feedback_id: int,
    db: Session = Depends(get_db),
    runtime: RuntimeSession = Depends(require_runtime_session),
) -> None:
    feedback = db.get(Feedback, feedback_id)
    if feedback is None or feedback.deleted_at is not None:
        raise HTTPException(status_code=404, detail="反馈不存在")
    feedback.deleted_at = datetime.now(timezone.utc)
    feedback.deleted_by_name = runtime.operator_name
    db.commit()
